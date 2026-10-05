import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def exportar_asistencias_a_excel(registros, ruta_salida=None):
    if ruta_salida is None:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        ruta_salida = f'bitacora_asistencias_{timestamp}.xlsx'

    wb = Workbook()
    ws = wb.active
    ws.title = "Bitácora de Asistencias"

    encabezado_font = Font(bold=True, color="FFFFFF", size=12)
    encabezado_fill = PatternFill(start_color="2E5E4E", end_color="2E5E4E", fill_type="solid")
    encabezado_align = Alignment(horizontal="center", vertical="center")

    borde = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    encabezados = [
        'Matrícula',
        'Apellido Paterno',
        'Apellido Materno',
        'Nombre',
        'Grupo',
        'Materia',
        'Profesor',
        'Fecha',
        'Hora',
        'Tipo',
        'Método',
        'Confianza',
        'Observaciones',
    ]

    for col_num, encabezado in enumerate(encabezados, 1):
        celda = ws.cell(row=1, column=col_num, value=encabezado)
        celda.font = encabezado_font
        celda.fill = encabezado_fill
        celda.alignment = encabezado_align
        celda.border = borde

    for fila_num, registro in enumerate(registros, 2):
        alumno = registro.alumno
        grupo = alumno.grupo.nombre if alumno.grupo else "Sin grupo"
        materia = registro.materia.nombre if registro.materia else "—"
        profesor = str(registro.profesor) if registro.profesor else "—"

        datos = [
            alumno.matricula,
            alumno.apellido_paterno,
            alumno.apellido_materno,
            alumno.nombre,
            grupo,
            materia,
            profesor,
            registro.fecha.strftime('%d/%m/%Y') if registro.fecha else '',
            registro.hora.strftime('%H:%M:%S') if registro.hora else '',
            registro.get_tipo_display(),
            registro.get_metodo_display(),
            f"{registro.confianza:.2f}" if registro.confianza is not None else '',
            registro.observaciones or '',
        ]

        for col_num, valor in enumerate(datos, 1):
            celda = ws.cell(row=fila_numa, column=col_num, value=valor) if False else ws.cell(row=fila_num, column=col_num, value=valor)
            celda.border = borde
            celda.alignment = Alignment(horizontal="center", vertical="center")

    anchos = [12, 18, 18, 18, 10, 18, 25, 12, 12, 12, 14, 10, 30]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    ws.freeze_panes = "A2"
    wb.save(ruta_salida)
    return ruta_salida